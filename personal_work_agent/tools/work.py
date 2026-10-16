from pathlib import Path
from datetime import datetime
import platform
import psutil
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from config import WORKSPACE

def safe_path(path: str) -> Path:
    WORKSPACE.mkdir(parents=True, exist_ok=True)
    p = (WORKSPACE / path).resolve()
    if not p.is_relative_to(WORKSPACE.resolve()):
        raise ValueError('Đường dẫn nằm ngoài workspace được phép')
    return p

def list_files(folder: str = '.') -> str:
    p = safe_path(folder)
    if not p.is_dir(): return 'Không phải thư mục hợp lệ'
    return '\n'.join(str(x.relative_to(WORKSPACE)) for x in sorted(p.iterdir())[:100]) or 'Thư mục trống'

def read_document(path: str) -> str:
    p = safe_path(path)
    if not p.is_file(): return 'Không tìm thấy file'
    if p.stat().st_size > 10_000_000: return 'File quá lớn (giới hạn 10 MB)'
    ext = p.suffix.lower()
    if ext in ('.txt', '.md', '.py', '.csv', '.json'):
        content = p.read_text(encoding='utf-8-sig')
    elif ext == '.pdf':
        content = '\n'.join(page.extract_text() or '' for page in PdfReader(p).pages[:30])
    elif ext == '.docx':
        content = '\n'.join(x.text for x in Document(p).paragraphs)
    elif ext == '.xlsx':
        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            content = '\n'.join(f'{ws.title}: ' + ' | '.join(' ; '.join(str(v) for v in row if v is not None) for row in ws.iter_rows(max_row=100, max_col=20, values_only=True)) for ws in wb.worksheets[:5])
        finally: wb.close()
    else: return 'Chưa hỗ trợ định dạng này'
    return content[:16000] or 'Không trích xuất được văn bản (PDF scan có thể cần OCR)'

def system_info() -> str:
    mem = psutil.virtual_memory()
    return f'Windows/OS: {platform.platform()} | RAM: {mem.percent}% sử dụng | CPU: {psutil.cpu_percent(interval=0.1)}%'

def current_time() -> str:
    return datetime.now().strftime('%d/%m/%Y %H:%M:%S')
